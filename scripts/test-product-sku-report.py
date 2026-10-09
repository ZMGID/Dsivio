"""Exercise both packaged report entry points with links and long opportunity IDs.

Run with the report requirements installed:
  uv run --with-requirements src-tauri/resources/plugins/shopee-research/skills/shopee-research/scripts/requirements.txt python scripts/test-product-sku-report.py
"""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]


class ReportTest(unittest.TestCase):
    def test_packaged_reports_preserve_lookup_identifiers(self):
        for plugin in ('shopee-research', 'product-sku-search'):
            for identifiers in (['url'], ['opportunityId'], ['url', 'opportunityId']):
                with self.subTest(plugin=plugin, identifiers=identifiers), tempfile.TemporaryDirectory() as tmp:
                    directory = Path(tmp)
                    Image.new('RGB', (100, 100), 'white').save(directory / 'product.png')
                    columns = ['name', 'image', 'package', 'price', *identifiers]
                    candidate = dict(id='test', name='Product', unit='件', salesUnit='件', qty=1,
                                     price=19.9, image='product.png',
                                     opportunityId='12345678901234567890', url='https://example.com/product')
                    for key in ('url', 'opportunityId'):
                        if key not in identifiers:
                            del candidate[key]
                    data = dict(title='Report', summary='Summary', analysis='Analysis', currency='BRL',
                                columns=columns, candidates=[candidate])
                    source = directory / 'report.json'
                    source.write_text(json.dumps(data), encoding='utf-8')
                    output = directory / 'report.xlsx'
                    script = ROOT / f'src-tauri/resources/plugins/{plugin}/skills/{plugin}/scripts/build_report.py'
                    result = subprocess.run([sys.executable, str(script), str(source), str(output)], capture_output=True, text=True)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    wb = load_workbook(output)
                    sheet = wb.active
                    for key in identifiers:
                        cell = sheet.cell(7, columns.index(key) + 1)
                        self.assertEqual(cell.value, candidate[key])
                        self.assertEqual(cell.data_type, 's')
                        if key == 'url':
                            self.assertEqual(cell.hyperlink.target, candidate[key])
                    self.assertEqual(len(sheet._images), 1)
                    wb.close()


if __name__ == '__main__':
    unittest.main()

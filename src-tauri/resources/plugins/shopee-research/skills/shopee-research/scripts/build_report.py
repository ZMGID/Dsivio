# coding: utf-8
"""Generate a compact research workbook using public Python dependencies."""
import argparse
import base64
from datetime import datetime, timezone
from html import escape
import json
import math
import os
from pathlib import Path
import tempfile
import time
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from finalize_report import finalize

FIELDS = {
    'name': ('编号／商品名称', 180), 'image': ('图片', 160), 'size': ('产品规格（CM）', 170),
    'material': ('材质', 185), 'features': ('类型／功能', 210), 'package': ('套装内容／销售单位', 200),
    'color': ('颜色／款式', 180), 'price': ('竞品价格', 135), 'url': ('竞品链接', 240), 'sales': ('销量', 145),
}
DEFAULT_COLUMNS = ['name', 'image', 'size', 'material', 'package', 'color', 'price', 'url', 'sales']


def layout(data):
    keys = data.get('columns', DEFAULT_COLUMNS)
    if (not isinstance(keys, list) or any(not isinstance(k, str) or k not in FIELDS for k in keys)
            or len(set(keys)) != len(keys)
            or not set(('name', 'image', 'package', 'price', 'url', 'sales')).issubset(keys)):
        raise ValueError('Invalid report columns')
    return [(k, FIELDS[k][0] + (f'（{data["currency"]}）' if k == 'price' else
                              f'（{data["salesPeriod"]}）' if k == 'sales' else ''), FIELDS[k][1]) for k in keys]


def validate(data):
    for key in ('title', 'summary', 'analysis', 'currency', 'salesPeriod'):
        if not isinstance(data.get(key), str) or not data[key].strip():
            raise ValueError(f'Missing {key}')
    if len(data['currency']) != 3 or not data['currency'].isascii() or not data['currency'].isalpha() or not data['currency'].isupper():
        raise ValueError('Use an ISO currency code')
    layout(data)
    if not isinstance(data.get('candidates'), list) or not data['candidates']:
        raise ValueError('No candidates')
    seen = set()
    for c in data['candidates']:
        for key in ('id', 'name', 'unit', 'salesUnit', 'url', 'image'):
            if not isinstance(c.get(key), str) or not c[key].strip():
                raise ValueError(f'Missing candidate {key}')
        if not c['id'].isascii() or not all(ch.isalnum() or ch in '_-' for ch in c['id']) or c['id'] in seen:
            raise ValueError('Invalid/duplicate ID')
        seen.add(c['id'])
        if (type(c.get('price')) not in (int, float) or not math.isfinite(c['price']) or c['price'] <= 0
                or type(c.get('qty')) is not int or c['qty'] <= 0):
            raise ValueError('Invalid price/quantity')
        sales = c.get('sales')
        if not ((type(sales) in (int, float) and math.isfinite(sales) and sales >= 0)
                or (isinstance(sales, str) and sales.strip())):
            raise ValueError('Missing sales')
        if urlparse(c['url']).scheme not in ('http', 'https') or not urlparse(c['url']).netloc:
            raise ValueError('Invalid URL')
    return data


def values(c, index):
    return dict(c, name=f'H{index+1} {c["name"]}' + ('\n'+c['role'] if c.get('role') else ''), image=None,
                package='\n'.join([f'{c["qty"]}{c["unit"]}／{c["salesUnit"]}'] +
                                  [c[k] for k in ('packContents', 'packSize', 'quantityBasis') if c.get(k)]))


def preview(data, columns, work):
    """Self-contained layout preview, with the same text and compressed pictures."""
    parts = ['<!doctype html><meta charset="utf-8"><title>调研表布局预览</title>',
             '<style>body{font:14px Arial,sans-serif}table{border-collapse:collapse;table-layout:fixed}'
             'td,th{border:1px solid #333;white-space:pre-wrap;overflow-wrap:anywhere;text-align:center;'
             'vertical-align:middle;padding:2px;box-sizing:border-box}td{height:160px}img{width:156px;height:156px;'
             'object-fit:contain}a{color:#1762a6}header{max-width:'+str(sum(c[2] for c in columns))+'px}</style><header>']
    for tag, key in [('h2', 'title'), ('p', 'summary'), ('p', 'analysis')]:
        parts.append(f'<{tag}>{escape(data[key])}</{tag}>')
    parts.append('</header><table><colgroup>')
    parts.extend(f'<col style="width:{width}px">' for _, _, width in columns)
    parts.append('</colgroup><tr>')
    parts.extend('<th>'+escape(label)+'</th>' for _, label, _ in columns)
    parts.append('</tr>')
    for i, c in enumerate(data['candidates'][:5]):
        row = values(c, i)
        parts.append('<tr>')
        for key, _, _ in columns:
            if key == 'image':
                content = '<img alt="'+escape(c['name'], quote=True)+'" src="data:image/jpeg;base64,'+base64.b64encode((work/(c['id']+'.jpg')).read_bytes()).decode()+'">'
            elif key == 'url':
                content = '<a href="'+escape(c['url'], quote=True)+'">'+escape(c['url'])+'</a>'
            else:
                content = escape(f'{data["currency"]} {c["price"]:.2f}' if key == 'price' else str(row.get(key) if row.get(key) is not None else ''))
            parts.append('<td>'+content+'</td>')
        parts.append('</tr>')
    return ''.join(parts)+'</table>'


def build(input_path, output):
    started = time.perf_counter()
    input_path, output = Path(input_path).resolve(), Path(output).resolve()
    data = validate(json.loads(input_path.read_text(encoding='utf-8')))
    columns = layout(data)
    positions = {key: get_column_letter(i+1) for i, (key, _, _) in enumerate(columns)}
    progress_path = input_path.parent/'progress.json'
    progress = json.loads(progress_path.read_text(encoding='utf-8')) if progress_path.exists() else {}
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.research-', dir=output.parent) as temporary:
        work = Path(temporary)
        wb = Workbook()
        sheet = wb.active
        sheet.title = '竞品调研表'
        sheet.sheet_view.showGridLines = False
        for i, (_, _, width) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = (width-5)/7
        last = len(columns)
        for r, key in enumerate(('title', 'summary', 'analysis'), 1):
            sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
            cell = sheet.cell(r, 1, data[key])
            cell.data_type = 's'
            cell.font = Font(name='Arial', size=16 if r == 1 else 11, bold=r == 1)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.row_dimensions[r].height = (34 if r == 1 else max(28 if r == 2 else 44, math.ceil(len(data[key])/100)*22))*.75
        border = Border(**{side: Side(style='thin', color='333333') for side in ('top', 'bottom', 'left', 'right')})
        rows = [[label for _, label, _ in columns]]
        rows += [[values(c, i).get(key) for key, _, _ in columns] for i, c in enumerate(data['candidates'])]
        for r, row in enumerate(rows, 6):
            sheet.row_dimensions[r].height = 33.75 if r == 6 else 120
            for col, value in enumerate(row, 1):
                cell = sheet.cell(r, col, value)
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Arial', size=11, bold=r == 6, color='1762A6' if r > 6 and columns[col-1][0] == 'url' else '000000')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
            if r > 6:
                sheet[f'{positions["price"]}{r}'].number_format = f'"{data["currency"]} "0.00'
        if len(data['candidates']) > 5:
            sheet.freeze_panes = 'A7'
        wb.save(work/'table-base.xlsx')
        prepared = time.perf_counter()
        result = finalize(input_path, work/'table-base.xlsx', work/'final.xlsx', positions)
        (work/'report-layout.html').write_text(preview(data, columns, work), encoding='utf-8')
        timings = dict(tableSeconds=prepared-started, imagesAndChecksSeconds=time.perf_counter()-prepared)
        progress['report'] = dict(status='complete', output=str(output), checkedAt=datetime.now(timezone.utc).isoformat(), **result, timings=timings)
        (work/'progress.json').write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(work/'final.xlsx', output)
        os.replace(work/'report-layout.html', input_path.parent/'report-layout.html')
        os.replace(work/'progress.json', progress_path)
    return dict(output=str(output), **result, timings=timings)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try:
        print(json.dumps(build(args.input, args.output), ensure_ascii=False))
    except Exception as error:
        parser.exit(1, f'{error}\n')

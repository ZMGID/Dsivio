# coding: utf-8
import copy
import json
from pathlib import Path
import tempfile
import subprocess
import sys
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image
from build_report import build, validate
from embed_images import S, X, xml
from setup_runtime import setup


def fixture():
    return dict(title='测试调研', summary='三个样本，按页面销量量级筛选。', analysis='七双装折合BRL 1.50/双。',
                currency='BRL', salesPeriod='累计', candidates=[dict(id='123', name='花纹袜', qty=7, unit='双',
                salesUnit='套', price=10.49, sales='40mil+', url='https://shopee.com.br/product/1/123', image='123.png')])


class BuilderTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.input, self.output = self.root/'report.json', self.root/'report.xlsx'
        Image.new('RGB', (700, 900), 'lightblue').save(self.root/'123.png')

    def run_build(self, data):
        self.input.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        return build(self.input, self.output)

    def test_validation_before_export(self):
        validate(fixture())
        for key in ('qty', 'price', 'unit', 'salesUnit', 'image'):
            data = fixture()
            del data['candidates'][0][key]
            with self.subTest(key=key), self.assertRaises(ValueError):
                validate(data)
        for invalid in (['name', 'image'], ['name', 'image', 'package', 'price', 'url', 'sales', 'image']):
            data = fixture()
            data['columns'] = invalid
            with self.assertRaises(ValueError):
                validate(data)
        data = fixture()
        data['candidates'].append(copy.deepcopy(data['candidates'][0]))
        with self.assertRaises(ValueError):
            validate(data)

    def test_pipeline_dynamic_columns_progress_and_failure_preserves_output(self):
        data = fixture()
        data['columns'] = ['name', 'features', 'image', 'package', 'color', 'price', 'url', 'sales']
        data['candidates'] = [dict(data['candidates'][0], id=str(123+i), features='可躺；便携折叠', packContents='单车',
                                   url=f'https://shopee.com.br/product/1/{123+i}') for i in range(3)]
        (self.root/'progress.json').write_text(json.dumps(dict(step='出表', notes=['保留采集记录'])))
        result = self.run_build(data)
        self.assertTrue(result['fileChecks'])
        self.assertEqual(result['pictureCount'], 3)
        workbook = load_workbook(self.output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ['竞品调研表'])
        sheet = workbook.active
        self.assertEqual(sheet['B6'].value, '类型／功能')
        self.assertEqual(sheet['B7'].value, '可躺；便携折叠')
        self.assertEqual(sheet['F7'].value, 10.49)
        self.assertEqual(sheet['G7'].hyperlink.target, data['candidates'][0]['url'])
        self.assertEqual(len(sheet._images), 3)
        self.assertTrue(sheet.protection.objects)
        self.assertFalse(sheet['A7'].protection.locked)
        progress = json.loads((self.root/'progress.json').read_text())
        self.assertEqual(progress['notes'], ['保留采集记录'])
        self.assertEqual(progress['report']['status'], 'complete')
        preview = (self.root/'report-layout.html').read_text(encoding='utf-8')
        self.assertIn('可躺；便携折叠', preview)
        self.assertEqual(preview.count('data:image/jpeg;base64,'), 3)
        with ZipFile(self.output) as archive:
            self.assertEqual(len([n for n in archive.namelist() if n.startswith('xl/media/')]), 1)
            root = xml(archive.read('xl/worksheets/sheet1.xml'))
            names = [n.tag.split('}')[-1] for n in root]
            self.assertLess(names.index('hyperlinks'), names.index('pageMargins'))
            drawing = xml(archive.read('xl/drawings/research1.xml'))
            self.assertEqual(drawing[0].find(f'{{{X}}}from/{{{X}}}col').text, '2')
        original = self.output.read_bytes()
        data['candidates'][0]['image'] = 'missing.png'
        with self.assertRaises(FileNotFoundError):
            self.run_build(data)
        self.assertEqual(self.output.read_bytes(), original)
        self.assertFalse(list(self.root.glob('.research-*')))

    def test_forty_rows_keep_compact_images_and_literal_text(self):
        data = fixture()
        data['candidates'] = [dict(data['candidates'][0], id=str(i), name='=1+1',
                                   url=f'https://shopee.com.br/product/1/{i}') for i in range(40)]
        self.input.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        run = subprocess.run([sys.executable, '-B', str(Path(__file__).with_name('build_report.py')),
                              str(self.input), str(self.output)], check=True, capture_output=True, text=True)
        result = json.loads(run.stdout)
        self.assertEqual(result['pictureCount'], 40)
        self.assertLess(result['imageBytes'], 40960)
        self.assertLess(self.output.stat().st_size, 100000)
        workbook = load_workbook(self.output)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, 46)
        self.assertEqual(workbook.active['A7'].data_type, 's')
        self.assertEqual(workbook.active.freeze_panes, 'A7')

    def test_runtime_check_reports_missing_without_creating_environment(self):
        target = self.root/'runtime'
        self.assertFalse(setup(target, check=True)['ready'])
        self.assertFalse(target.exists())


if __name__ == '__main__':
    unittest.main()

"""Load editable report state, freeze run configuration, and validate collected data."""
from __future__ import annotations

import argparse
import copy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    temporary.replace(path)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def safe_name(value):
    return isinstance(value, str) and value.strip() and not re.search(r'[<>:"/\\|?*\x00-\x1f]', value) and value not in ('.', '..')


def report_outputs(report):
    if 'outputs' in report:
        return report['outputs']
    # Read old saved run snapshots without changing their original output paths.
    result = [{'id': 'brief', 'template': report.get('template'), 'renderer': 'summary',
               'formats': ['html', 'png'], 'directory': '日报'}]
    if report.get('operationsTemplate'):
        result.append({'id': 'operations', 'template': report['operationsTemplate'], 'renderer': 'operations',
                       'formats': ['html'], 'directory': '运营分析', 'suffix': '运营分析'})
    return result


def output_paths(report, day):
    definitions = report_outputs(report)
    require(isinstance(definitions, list) and definitions, 'outputs 至少配置一份报告')
    result, used_paths = {}, set()
    for item in definitions:
        require(safe_name(item.get('id')) and item['id'] not in result, '报告 id 无效或重复')
        require(item.get('renderer') in ('summary', 'operations'), '报告 renderer 必须为 summary 或 operations')
        require(safe_name(item.get('template')) and (ROOT / 'templates' / f"{item['template']}.html").is_file(), '所选报告模板不存在')
        marker = '__DATA__' if item['renderer'] == 'operations' else '__REPORT_JSON__'
        require(marker in (ROOT / 'templates' / f"{item['template']}.html").read_text(encoding='utf-8'), '模板与所选渲染器不匹配')
        formats = item.get('formats', ['html'])
        require(formats in (['html'], ['html', 'png']), '报告格式支持 html 或 html+png')
        require(item['renderer'] != 'operations' or formats == ['html'], '运营报告仅生成 HTML')
        directory = item.get('directory', item['id'])
        require(safe_name(directory), '报告目录名无效')
        suffix = item.get('suffix', '')
        require(not suffix or safe_name(suffix), '报告文件后缀无效')
        base = f"{report['filePrefix']}_{day}" + (f'_{suffix}' if suffix else '')
        folder = Path(report['outputRoot']) / directory / day
        files = {kind: str(folder / f'{base}.{kind}') for kind in formats}
        normalized = {str(Path(p).resolve()).casefold() for p in files.values()}
        require(not used_paths.intersection(normalized), '不同报告的保存路径不能重复')
        used_paths.update(normalized)
        result[item['id']] = files
    return result


def resolve_config(config_path=None, report_id=None):
    require(not (config_path and report_id), '--config 与 --report 不能同时指定')
    if config_path:
        return Path(config_path)
    if report_id:
        require(safe_name(report_id), '日报配置名称无效')
        return ROOT / 'reports' / f'{report_id}.json'
    candidates = sorted((ROOT / 'reports').glob('*.json'))
    require(len(candidates) == 1, '存在多份日报或没有日报，请用 --report 日报名 或 --config 文件路径明确指定')
    return candidates[0]


def load_snapshot(state_path=None, report_id=None):
    config_path = resolve_config(state_path, report_id)
    state = read_json(config_path)
    require(state.get('schemaVersion') == 1, '日报配置 schemaVersion 必须为 1')
    report = state['report']
    require(Path(report['outputRoot']).is_absolute(), 'report.outputRoot 必须是绝对路径')
    for key in ('filePrefix', 'workbookName'):
        require(safe_name(report.get(key)), f'report.{key} 必须是有效文件名')
    require(report['workbookName'].endswith('.xlsx'), 'workbookName 必须为 .xlsx')
    require(isinstance(report.get('title'), str) and report['title'].strip(), '缺少报告名称')
    output_paths(report, '2000-01-01')
    require(report.get('delivery') in ('png_only', 'local_links'), 'delivery 必须为 png_only 或 local_links')
    require(isinstance(report.get('dateUtcOffsetMinutes'), int) and -720 <= report['dateUtcOffsetMinutes'] <= 840, '默认日期时区偏移无效')
    require(isinstance(state.get('requirements'), list) and all(isinstance(x, str) for x in state['requirements']), 'requirements 必须是文本列表')
    registry = read_json(ROOT / 'platforms/registry.json')
    stores = []
    ids = set()
    for item in state['stores']:
        require(isinstance(item.get('enabled'), bool), '每家店必须配置 enabled')
        require(isinstance(item.get('storeId'), str) and item['storeId'].isdigit(), 'storeId 必须是数字字符串')
        require(item['storeId'] not in ids, 'storeId 重复')
        ids.add(item['storeId'])
        if not item['enabled']:
            continue
        platform = registry.get(item['platform'])
        require(platform is not None, f"平台未注册: {item['platform']}")
        require((ROOT / 'platforms' / platform['guide']).is_file(), '平台文档不存在')
        require(item['region'] in platform['regions'], f"当前平台适配未验证该地区: {item['region']}")
        require(re.fullmatch(r'[A-Z]{3}', item.get('currency', '')) is not None, 'currency 必须是三位币种代码')
        require(all(isinstance(item.get(k), str) and item[k].strip() for k in ('name', 'expectedName', 'timeZone')), '缺少店铺名称或时区')
        require(isinstance(item.get('collectAds'), bool), 'collectAds 必须是布尔值')
        require(not item['collectAds'] or platform['supportsAds'], '该平台尚未支持广告取数')
        require(item['orderMetric'] in platform['orderMetrics'], '平台不支持所选订单口径')
        require(isinstance(item.get('requirements'), list) and all(isinstance(x, str) for x in item['requirements']), '店铺 requirements 必须是文本列表')
        if item.get('ledgerPath'):
            require(Path(item['ledgerPath']).is_absolute() and item['ledgerPath'].endswith('.xlsx'), 'ledgerPath 必须为绝对 xlsx 路径')
        stores.append({**copy.deepcopy(item), 'platformName': platform['name'], 'orderLabel': platform['orderMetrics'][item['orderMetric']], 'guide': platform['guide']})
    require(stores, '没有启用的店铺，不能生成空日报')
    snapshot = {'configFile': str(config_path.resolve()), 'report': copy.deepcopy(report), 'requirements': copy.deepcopy(state['requirements']), 'stores': stores}
    snapshot['configHash'] = hashlib.sha256(json.dumps(snapshot, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return snapshot


def target_date(snapshot, requested=None):
    if requested:
        return date.fromisoformat(requested).isoformat()
    offset = snapshot['report']['dateUtcOffsetMinutes']
    return (datetime.now(timezone(timedelta(minutes=offset))).date() - timedelta(days=1)).isoformat()


def paths(snapshot, day):
    day = date.fromisoformat(day).isoformat()
    root = Path(snapshot['report']['outputRoot'])
    base = f"{snapshot['report']['filePrefix']}_{day}"
    if 'outputs' in snapshot['report']:
        return {'json': str(root / '数据' / day / f'report-data_{day}.json'),
                'workbook': str(root / snapshot['report']['workbookName']),
                'reports': output_paths(snapshot['report'], day)}
    output = {'json': str(root / '数据' / day / f'report-data_{day}.json'),
            'html': str(root / '日报' / day / f'{base}.html'),
            'png': str(root / '日报' / day / f'{base}.png'),
            'workbook': str(root / snapshot['report']['workbookName'])}
    if snapshot['report'].get('operationsTemplate'):
        output['operationsHtml'] = str(root / '运营分析' / day / f'{base}_运营分析.html')
    return output


def number(value, label, integer=False):
    require(not isinstance(value, bool) and value is not None, f'{label} 缺失或类型无效')
    try:
        result = Decimal(str(value).replace(',', '.'))
    except InvalidOperation:
        raise ValueError(f'{label} 不是数值') from None
    require(result.is_finite() and result >= 0, f'{label} 必须是有限非负数')
    if integer:
        require(result == result.to_integral_value(), f'{label} 必须是整数')
        return int(result)
    require(result == result.quantize(Decimal('.01')), f'{label} 最多保留两位小数')
    return f'{result:.2f}'


def prepare_report(report, snapshot=None):
    snapshot = copy.deepcopy(snapshot or report.get('configSnapshot') or load_snapshot())
    require(re.fullmatch(r'\d{4}-\d{2}-\d{2}', report['dateIso']) is not None, 'dateIso 必须为 YYYY-MM-DD')
    date.fromisoformat(report['dateIso'])
    incoming = report['stores']
    require(len(incoming) == len(snapshot['stores']), '结果数量必须与本次启用店铺一致；失败店也必须保留')
    used = set()
    normalized = []
    for config in snapshot['stores']:
        matches = [(index, row) for index, row in enumerate(incoming)
                   if row.get('storeId') == config['storeId'] or
                   (not row.get('storeId') and row.get('name') == config['name'])]
        require(len(matches) == 1 and matches[0][0] not in used, f"店铺结果缺失、重复或身份有歧义: {config['name']}")
        index, raw = matches[0]
        used.add(index)
        row = copy.deepcopy(raw)
        for key in ('storeId', 'name', 'platform', 'platformName', 'currency', 'orderMetric', 'orderLabel', 'collectAds'):
            if key in raw and key not in ('name', 'platformName', 'orderLabel'):
                require(raw[key] == config[key], f"{config['name']} 的 {key} 与配置不一致")
            row[key] = config[key]
        for fields, reason in [(('sales', 'orders'), 'missingReason'), (('adCost', 'roi'), 'adMissingReason')]:
            if reason == 'adMissingReason' and not config['collectAds']:
                row.update(adCost=None, roi=None)
                row.pop(reason, None)
                continue
            # Legacy whole-store failures may omit the separate advertising reason.
            if reason == 'adMissingReason' and row.get('missingReason') and all(row.get(k) is None for k in fields):
                row.setdefault(reason, row['missingReason'])
            missing = [key for key in fields if row.get(key) is None]
            require(not missing or row.get(reason), f"{','.join(missing)} 缺失时须说明原因")
            require(missing or not row.get(reason), f'{reason} 与已完整获取的指标冲突')
            for key in fields:
                row[key] = None if row.get(key) is None else number(row[key], key, integer=key == 'orders')
        required = ['sales', 'orders'] + (['adCost', 'roi'] if config['collectAds'] else [])
        present = sum(row.get(key) is not None for key in required)
        row['status'] = '成功' if present == len(required) else '未获取' if present == 0 else '部分获取'
        normalized.append(row)
    return {**copy.deepcopy(report), 'configSnapshot': snapshot, 'stores': normalized}


EXPORT_TYPES = {
    'shopee': {'店铺经营'},
    'tiktok': {'关键指标'},
    'shein': {'GMV', '销量', '支付订单数'},
}


def _same_number(left, right, integer=False):
    if left is None or right is None:
        return left is right
    return number(left, '核验值', integer=integer) == number(right, '核验值', integer=integer)


def _target_record(store, report_day):
    records = [row for row in store.get('dailyRecords', []) if row.get('dateIso') == report_day]
    require(len(records) == 1, f"{store['name']} 必须恰好保留一条目标日每日记录: {report_day}")
    return records[0]


def _check_target_values(store, report_day):
    target = _target_record(store, report_day)
    platform = store['platform']
    if platform == 'shopee':
        expected_sales, expected_orders = target.get('sales'), target.get('orders')
    elif platform == 'tiktok':
        expected_sales = target.get('sales', target.get('GMV'))
        expected_orders = target.get('orders', target.get('订单数'))
        if store.get('itemsSold') is not None:
            expected_items = target.get('itemsSold', target.get('unitsSold', target.get('商品成交件数')))
            require(_same_number(store['itemsSold'], expected_items, integer=True),
                    f"{store['name']} 的商品成交件数与目标日 Excel 行不一致")
    else:
        expected_sales = target.get('sales', target.get('GMV'))
        expected_orders = target.get('unitsSold', target.get('销量')) if store['orderMetric'] == 'unitsSold' else target.get('paidOrders', target.get('支付订单数'))
        if store.get('paidOrders') is not None:
            expected_paid = target.get('paidOrders', target.get('支付订单数'))
            require(_same_number(store['paidOrders'], expected_paid, integer=True),
                    f"{store['name']} 的支付订单数与目标日 Excel 行不一致")
    require(_same_number(store.get('sales'), expected_sales), f"{store['name']} 的销售额与目标日 Excel 行不一致")
    require(_same_number(store.get('orders'), expected_orders, integer=True), f"{store['name']} 的数量与目标日 Excel 行不一致")
    if store.get('source') == 'single-day-page-realtime-updating':
        require(platform == 'shopee' and target.get('source') == 'single-day-page-realtime-updating',
                f"{store['name']} 的更新中网页覆盖证据不完整")


def _manifest_entries(report):
    from download_files import check_downloads, digest, inside
    check_downloads(report)
    root = Path(report['configSnapshot']['report']['outputRoot']).resolve()
    stores = {store['storeId'] for store in report['stores']}
    result = []
    for filename in report['downloads']:
        manifest = Path(filename)
        entry = read_json(manifest)
        require(entry.get('storeId') in stores, f'下载清单店铺不属于本次报告: {manifest}')
        require(entry.get('type') in {'店铺经营', '广告组数据', '关键指标', 'GMV', '销量', '支付订单数'},
                f'下载清单数据类型无效: {manifest}')
        start, end = date.fromisoformat(entry['start']), date.fromisoformat(entry['end'])
        report_day = date.fromisoformat(report['dateIso'])
        require(start <= report_day <= end, f'下载清单未覆盖目标日 {report["dateIso"]}: {manifest}')
        archive = Path(entry['path'])
        require(inside(archive, root / '数据') and archive.is_file(), f'归档文件缺失: {archive}')
        require(entry.get('sha256') == digest(archive), f'归档文件哈希不一致: {archive}')
        result.append(entry)
    return result


def _check_central_ledger(report, workbook):
    from openpyxl import load_workbook
    book = load_workbook(workbook, data_only=True, read_only=True)
    try:
        sheet = book['日报明细']
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        for store in report['stores']:
            matches = [row for row in rows if str(row[0])[:10] == report['dateIso'] and str(row[12]) == store['storeId']]
            require(len(matches) == 1, f"累计台账缺少或重复目标日店铺行: {store['name']}")
            row = matches[0]
            require(_same_number(row[3], store.get('sales')), f"累计台账销售额不一致: {store['name']}")
            require(_same_number(row[4], store.get('orders'), integer=True), f"累计台账数量不一致: {store['name']}")
            require(row[7] == store['status'], f"累计台账状态不一致: {store['name']}")
    finally:
        book.close()


def _check_artifact(path, report_day):
    path = Path(path)
    require(path.is_file() and path.stat().st_size > 0, f'正式产物缺失或为空: {path}')
    if path.suffix.lower() == '.html':
        text = path.read_text(encoding='utf-8')
        require(report_day in text, f'HTML 未显示目标日期 {report_day}: {path}')
        require('[object Object]' not in text, f'HTML 含未正确渲染的数据对象: {path}')
    elif path.suffix.lower() == '.png':
        header = path.read_bytes()[:24]
        require(len(header) == 24 and header[:8] == b'\x89PNG\r\n\x1a\n', f'PNG 文件无效: {path}')
        width, height = int.from_bytes(header[16:20], 'big'), int.from_bytes(header[20:24], 'big')
        require(width > 0 and height > 0, f'PNG 尺寸无效: {path}')


def audit_report(report, report_json):
    output = paths(report['configSnapshot'], report['dateIso'])
    require(Path(report_json).resolve() == Path(output['json']).resolve(), '正式 JSON 路径与冻结配置或业务日期不一致')
    entries = _manifest_entries(report)
    by_store = {}
    for entry in entries:
        by_store.setdefault(entry['storeId'], set()).add(entry['type'])
    for store in report['stores']:
        if store.get('sales') is None and store.get('orders') is None:
            continue
        _check_target_values(store, report['dateIso'])
        required = set(EXPORT_TYPES[store['platform']])
        if store.get('adCost') is not None or store.get('roi') is not None:
            required.add('广告组数据')
        missing = required - by_store.get(store['storeId'], set())
        require(not missing, f"{store['name']} 缺少目标日 Excel 归档: {','.join(sorted(missing))}")
        if store['platform'] == 'tiktok' and store.get('adCost') is not None:
            from tiktok_ads import parse as parse_ads
            ads = [e for e in entries if e['storeId'] == store['storeId'] and e['type'] == '广告组数据']
            require(len(ads) == 1 and ads[0]['start'] == ads[0]['end'] == report['dateIso'], 'TikTok 广告导出必须唯一且为目标单日')
            parsed = parse_ads(ads[0]['path'], report['dateIso'])
            require(_same_number(store['adCost'], parsed['adCost']) and _same_number(store['roi'], parsed['roi']), 'TikTok 广告成本/ROI与Excel不一致')
        captured = datetime.fromisoformat(store['capturedAt'])
        require(captured.tzinfo is not None, f"{store['name']} 的 capturedAt 缺少时区")
    workbook = Path(output['workbook'])
    require(workbook.is_file() and workbook.stat().st_size > 0, '累计 Excel 台账缺失或为空')
    _check_central_ledger(report, workbook)
    from store_ledger import check_ledgers
    check_ledgers(report)
    artifact_paths = [p for group in output_paths(report['configSnapshot']['report'], report['dateIso']).values() for p in group.values()]
    for artifact in artifact_paths:
        _check_artifact(artifact, report['dateIso'])
    return {'ok': True, 'dateIso': report['dateIso'], 'storeCount': len(report['stores']),
            'downloadCount': len(entries), 'artifacts': artifact_paths, 'workbook': str(workbook)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['plan', 'prepare', 'audit', 'complete'])
    parser.add_argument('report_json', nargs='?')
    parser.add_argument('--config', '--state', dest='config', help='独立日报配置文件路径')
    parser.add_argument('--report', help='reports 目录内的日报配置名称，不含 .json')
    parser.add_argument('--date')
    args = parser.parse_args()
    if args.action == 'plan':
        config_path = resolve_config(args.config, args.report)
        snapshot = load_snapshot(config_path)
        day = target_date(snapshot, args.date)
        print(json.dumps({'dateIso': day, 'configPath': str(config_path), 'configSnapshot': snapshot, 'paths': paths(snapshot, day)}, ensure_ascii=False, indent=2))
        return
    require(args.report_json, '必须提供结果 JSON 路径')
    original = read_json(args.report_json)
    if args.config or args.report:
        config_path = resolve_config(args.config, args.report)
    elif original.get('configSnapshot', {}).get('configFile'):
        config_path = Path(original['configSnapshot']['configFile'])
    else:
        config_path = resolve_config()
    report = prepare_report(original, original.get('configSnapshot') or load_snapshot(config_path))
    if args.action == 'prepare':
        write_json(args.report_json, report)
        print(json.dumps({'ok': True, 'storeCount': len(report['stores']), 'paths': paths(report['configSnapshot'], report['dateIso'])}, ensure_ascii=False))
        return
    output = paths(report['configSnapshot'], report['dateIso'])
    audit = audit_report(report, args.report_json)
    if args.action == 'audit':
        print(json.dumps(audit, ensure_ascii=False))
        return
    from download_files import cleanup
    source_config = report['configSnapshot'].get('configFile')
    if source_config:
        if args.config or args.report:
            require(config_path.resolve() == Path(source_config).resolve(), '不能把运行状态写入另一份日报配置')
        config_path = Path(source_config)
    state = read_json(config_path)
    cleanup_result = cleanup(report['configSnapshot']['report']['outputRoot'])
    state['lastRun'] = {'dateIso': report['dateIso'], 'finishedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
                        'status': '部分完成' if any(s['status'] != '成功' for s in report['stores']) else '完成',
                        'cleanupStatus': '待处理' if cleanup_result['pending'] else '完成',
                        'configHash': report['configSnapshot']['configHash'], 'paths': output,
                        'cleanup': cleanup_result,
                        'stores': [{'storeId': s['storeId'], 'status': s['status'],
                                    'reason': '；'.join(filter(None, [s.get('missingReason'), s.get('adMissingReason')]))} for s in report['stores']]}
    write_json(config_path, state)
    print(json.dumps({'ok': True, 'lastRun': state['lastRun']}, ensure_ascii=False))


if __name__ == '__main__':
    main()

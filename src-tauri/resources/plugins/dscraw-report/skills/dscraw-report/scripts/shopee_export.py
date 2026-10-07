"""Read Shopee BR exports without rewriting them or summing ad hierarchy rows."""
import csv
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
import re
from openpyxl import load_workbook


def number(value, brazil=True):
    if value is None or str(value).strip() in ('', '-', '--'):
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        result = Decimal(str(value))
    else:
        text = str(value).strip().replace('R$', '').strip()
        if brazil:
            if not re.fullmatch(r'-?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?', text):
                raise ValueError('巴西数字格式不明确')
            text = text.replace('.', '').replace(',', '.')
        elif not re.fullmatch(r'-?\d+(?:\.\d+)?', text):
            raise ValueError('CSV 数字格式不明确')
        result = Decimal(text)
    if not result.is_finite():
        raise ValueError('非有限数字')
    return result


def day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return datetime.strptime(str(value).strip(), '%d/%m/%Y').date().isoformat()


def business_rows(rows, start, end, target):
    header, daily, summary = None, {}, None
    date_key = sales_key = orders_key = None
    for rowno, row in enumerate(rows, 1):
        if not row or row[0] is None:
            continue
        if row[0] in ('日期', 'Data'):
            date_key = '日期' if row[0] == '日期' else 'Data'
            sales_key = '销售额 (BRL)' if '销售额 (BRL)' in row else 'Vendas (BRL)'
            orders_key = '订单数' if '订单数' in row else 'Pedidos'
            if any(row.count(key) != 1 for key in (date_key, sales_key, orders_key)):
                raise ValueError('经营表头缺失或重复')
            header = list(row)
            continue
        if header is None:
            continue
        if len(row) != len(header):
            raise ValueError('经营行列数不匹配')
        record = dict(zip(header, row))
        sales, orders = number(record[sales_key]), number(record[orders_key])
        if orders is not None and (orders < 0 or orders != orders.to_integral_value()):
            raise ValueError('订单数不是非负整数')
        item = {'sales': sales, 'orders': int(orders) if orders is not None else None,
                'row': rowno, 'raw': record}
        match = re.fullmatch(r'(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})', str(row[0]))
        if match:
            if summary is not None or (day(match[1]), day(match[2])) != (start, end):
                raise ValueError('汇总期间不匹配或重复')
            summary = item
        else:
            key = day(row[0])
            if key in daily:
                raise ValueError('重复日期')
            daily[key] = item
    expected = {(date.fromisoformat(start)+timedelta(days=i)).isoformat()
                for i in range((date.fromisoformat(end)-date.fromisoformat(start)).days+1)}
    if set(daily) != expected or target not in daily or summary is None:
        raise ValueError('日期不连续、不覆盖目标日或汇总缺失')
    totals = {key: sum(x[key] for x in daily.values()) if all(x[key] is not None for x in daily.values()) else None
              for key in ('sales', 'orders')}
    return {'sheet': '已下订单', 'start': start, 'end': end, 'daily': daily,
            'target': daily[target], 'summary': summary, 'totals': totals,
            'periodMismatch': any(totals[key] is None or totals[key] != summary[key] for key in totals)}


def parse_business(path, account, start, end, target):
    expected = account+'.shopee-shop-stats.'+start.replace('-', '')+'-'+end.replace('-', '')
    if not re.fullmatch(re.escape(expected)+r'(?:\s*\(\d+\))?\.xlsx', Path(path).name):
        raise ValueError('经营文件账号或期间不匹配')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = '已下订单' if '已下订单' in workbook.sheetnames else 'Pedido Feito'
        result = business_rows(list(workbook[sheet].values), start, end, target)
        result['sheet'] = sheet
        return result
    finally:
        workbook.close()


def parse_ads(path, account, target):
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.reader(handle))
    indices = [i for i, row in enumerate(rows) if '广告 / 商品名称' in row and '花费' in row]
    if len(indices) != 1:
        raise ValueError('广告表头缺失或重复')
    index = indices[0]
    metadata = {row[0]: row[1] for row in rows[:index] if len(row) == 2}
    dates = re.findall(r'\d{2}/\d{2}/\d{4}', metadata.get('时间', ''))
    if metadata.get('用户名称') != account or len(dates) != 2 or [day(x) for x in dates] != [target, target]:
        raise ValueError('广告 CSV 账号或日期不匹配')
    header, records = rows[index], []
    if len(set(header)) != len(header):
        raise ValueError('广告表头重复')
    for row in rows[index+1:]:
        if not row:
            continue
        if len(row) != len(header):
            raise ValueError('广告 CSV 行列数异常')
        record = dict(zip(header, row))
        for key in ('花费', '销售金额', '广告支出回报率'):
            number(record[key], brazil=False)
        records.append(record)
    return {'metadata': metadata, 'rowCount': len(records), 'rows': records,
            'coverageVerified': False, 'aggregationReason': '广告组与商品子行混排；未证明全店覆盖，不累加'}


def reconcile(parsed, single):
    reasons = {}
    result = {}
    used_realtime = False
    for key in ('sales', 'orders'):
        value = parsed['target'][key]
        reason = None
        if single['updating'] and value == 0:
            value = single[key]
            used_realtime = True
            if value is None:
                reason = '目标单日实时值缺失'
        elif value is None or (not single['updating'] and (single[key] is None or value != single[key])):
            reason = '目标日文件与单日页面缺失或不一致'
        result[key] = None if reason else value
        if reason:
            reasons[key] = reason
    return {**result, 'missingReason': reasons,
            'source': ('single-day-page-realtime-updating' if used_realtime
                       else 'export-and-single-day-page'),
            'periodMismatch': parsed['periodMismatch']}

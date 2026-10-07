"""Update each store's template workbook from already verified daily/ad records."""
from copy import copy
from datetime import date, datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from report_state import ROOT, number, require

DAILY = ['dateIso', 'sales', 'orders', 'unitsSold', 'adCost', 'adSales', 'roi',
         'conversions', 'cpa', 'status', 'reason', 'collectedAt', 'source', 'adSource']
ADS = ['dateIso', 'name', 'status', 'budget', 'targetRoas', 'adCost', 'adSales',
       'roi', 'conversions', 'cpa', 'reason', 'settingsCollectedAt', 'settingsDate',
       'collectedAt', 'adId', 'adType', 'metricBasis', 'dataStatus', 'source']
NUMERIC = {'sales', 'orders', 'unitsSold', 'adCost', 'adSales', 'roi', 'conversions',
           'cpa', 'budget', 'targetRoas'}
COUNTS = {'orders', 'unitsSold', 'conversions'}


def ledger_path(report, config):
    return Path(config.get('ledgerPath') or Path(report['configSnapshot']['report']['outputRoot']) /
                '店铺台账' / config['storeId'] / '经营台账.xlsx')


def daily_record(report, store):
    record = {key: store[key] for key in DAILY if key in store}
    record.update(dateIso=report['dateIso'], status='已核验' if store['status'] == '成功' else store['status'],
                  reason='；'.join(filter(None, [store.get('missingReason'), store.get('adMissingReason')])) or store.get('reason'),
                  collectedAt=store.get('collectedAt') or datetime.now().astimezone().isoformat(timespec='seconds'))
    if store['orderMetric'] == 'unitsSold':
        record['unitsSold'] = store['orders']
        record.pop('orders', None)
        if 'paidOrders' in store:
            record['orders'] = store['paidOrders']
    return record


def day(value):
    return value.date().isoformat() if isinstance(value, datetime) else str(value)[:10]


def cells(record, fields):
    values = []
    date.fromisoformat(record['dateIso'])
    for key in fields:
        value = record.get(key)
        if value is not None and key in NUMERIC:
            value = number(value, key, integer=key in COUNTS)
            value = int(value) if key in COUNTS else float(value)
        elif value and key in ('dateIso', 'settingsDate'):
            value = date.fromisoformat(value)
        values.append(value)
    return values


def upsert(sheet, records, fields, key_fields):
    indices = [fields.index(key) for key in key_fields]
    existing = {}
    for row in range(7, sheet.max_row + 1):
        if sheet.cell(row, 1).value is None:
            continue
        key = tuple(day(sheet.cell(row, i + 1).value) if i == 0 else sheet.cell(row, i + 1).value for i in indices)
        require(key not in existing, f'{sheet.title} 已有重复记录: {key}')
        existing[key] = row
    seen = set()
    for record in records:
        key = tuple(record.get(field) for field in key_fields)
        require(all(isinstance(item, str) and item.strip() for item in key), f'{sheet.title} 缺少日期或广告唯一标识')
        require(key not in seen, f'{sheet.title} 本次重复记录: {key}')
        seen.add(key)
        row = existing.get(key)
        if row is None:
            row = next((r for r in range(7, sheet.max_row + 1) if sheet.cell(r, 1).value is None), sheet.max_row + 1)
        for column, value in enumerate(cells(record, fields), 1):
            cell = sheet.cell(row, column)
            cell._style = copy(sheet.cell(7, column)._style)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
        existing[key] = row
    # Keep one chronological list, extending the existing table and dropdowns.
    ordered = sorted(([sheet.cell(r, c).value for c in range(1, len(fields) + 1)] for r in existing.values()),
                     key=lambda row: tuple(day(row[i]) if i == 0 else str(row[i]) for i in indices))
    for r, values in enumerate(ordered, 7):
        for c, value in enumerate(values, 1):
            sheet.cell(r, c).value = value
            if isinstance(value, str):
                sheet.cell(r, c).data_type = 's'
    end = max(16, 6 + len(ordered))
    for table in sheet.tables.values():
        table.ref = f'A6:{get_column_letter(len(fields))}{end}'
        if table.autoFilter:
            table.autoFilter.ref = table.ref
    for validation in sheet.data_validations.dataValidation:
        for area in list(validation.sqref.ranges):
            validation.sqref.remove(area)
            validation.add(f'{get_column_letter(area.min_col)}7:{get_column_letter(area.max_col)}{end}')


def update_ledgers(report):
    stamp = datetime.now().astimezone().strftime('%Y%m%d_%H%M%S_%f')
    for config, store in zip(report['configSnapshot']['stores'], report['stores']):
        path = ledger_path(report, config)
        book = load_workbook(path if path.exists() else ROOT / 'assets/store-ledger.xlsx')
        try:
            for sheet in book.worksheets:
                require(sheet['J2'].value in ('待填写', config['storeId']), '台账店铺身份不符')
                require(sheet['F2'].value in ('待填写', config['currency']), '台账币种不符')
                for cell, value in {'B2': config['name'], 'D2': config['platform'], 'F2': config['currency'],
                                    'H2': config['timeZone'], 'J2': config['storeId']}.items():
                    sheet[cell] = value
            current = daily_record(report, store)
            records = {r['dateIso']: r for r in store.get('dailyRecords', [])}
            require(len(records) == len(store.get('dailyRecords', [])), '本次每日记录日期重复')
            # The target date follows the report's verified values and failure state.
            records[report['dateIso']] = {**records.get(report['dateIso'], {}), **current}
            upsert(book['店铺日汇总'], records.values(), DAILY, ['dateIso'])
            upsert(book['广告明细'], store.get('adRecords', []) if config['collectAds'] else [], ADS,
                   ['dateIso', 'adType', 'adId', 'metricBasis'])
            path.parent.mkdir(parents=True, exist_ok=True)
            temporary = path.with_name(path.stem + '.tmp.xlsx')
            book.save(temporary)
            if path.exists():
                backup = path.parent / '历史版本' / stamp / path.name
                backup.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, backup)
            temporary.replace(path)
        finally:
            book.close()


def check_ledgers(report):
    for config, store in zip(report['configSnapshot']['stores'], report['stores']):
        path = ledger_path(report, config)
        require(path.is_file(), f'缺少每店台账: {config["name"]}')
        book = load_workbook(path)
        try:
            sheet = book['店铺日汇总']
            require(sheet['J2'].value == config['storeId'], '台账店铺身份不符')
            rows = [row for row in sheet.iter_rows(min_row=7, values_only=True) if row[0] and day(row[0]) == report['dateIso']]
            expected = cells(daily_record(report, store), DAILY)
            columns = [1, 3 if store['orderMetric'] == 'unitsSold' else 2, 4, 6, 9, 10]
            require(len(rows) == 1 and all(rows[0][i] == expected[i] for i in columns), f'每店台账目标日未更新: {config["name"]}')
        finally:
            book.close()

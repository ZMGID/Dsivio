#!/usr/bin/env python
"""Maintain the configured daily ledger, keyed by report date and store ID."""
from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from report_state import prepare_report, paths, read_json, write_json


HEADERS = ['日报日期', '店铺', '平台', '销售额', '数量', '广告花费', '广告ROI',
           '数据状态', '失败/缺失原因', '记录生成时间', '日报HTML', '日报PNG',
           '店铺ID', '币种', '数量口径']


def _new_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '日报明细'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    widths = [14, 20, 16, 16, 11, 18, 12, 14, 36, 24, 42, 42, 20, 10, 16]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.row_dimensions[1].height = 24
    return workbook


def _upgrade_legacy(sheet, configs):
    # Keep the existing twelve columns and historical business values in place.
    for column, title in enumerate(HEADERS, 1):
        if column > 12 and not sheet.cell(1, column).value:
            sheet.cell(1, column)._style = copy(sheet.cell(1, 3)._style)
            sheet.column_dimensions[get_column_letter(column)].width = {13:20, 14:10, 15:16}[column]
        sheet.cell(1, column).value = title
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 13).value:
            continue
        candidates = [s for s in configs if s['name'] == sheet.cell(row, 2).value
                      and s['platformName'] == sheet.cell(row, 3).value]
        if len(candidates) == 1:
            store = candidates[0]
            sheet.cell(row, 13).value = store['storeId']
            sheet.cell(row, 14).value = store['currency']
            sheet.cell(row, 15).value = store['orderMetric']


def update_ledger_and_enrich(report, workbook_path=None, output_dir=None, recorded_at=None):
    report = prepare_report(report)
    snapshot = report['configSnapshot']
    output = paths(snapshot, report['dateIso'])
    workbook_path = Path(workbook_path or output['workbook'])
    report_dir = Path(output_dir or snapshot['report']['outputRoot']) / '日报' / report['dateIso']
    base_name = f"{snapshot['report']['filePrefix']}_{report['dateIso']}"
    links = next(iter(output['reports'].values())) if 'reports' in output and not output_dir else {
        'html': str(report_dir / f'{base_name}.html'), 'png': str(report_dir / f'{base_name}.png')}
    workbook = load_workbook(workbook_path) if workbook_path.exists() else _new_workbook()
    sheet = workbook['日报明细']
    _upgrade_legacy(sheet, snapshot['stores'])
    existing = {}
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, 1).value
        day = value.date().isoformat() if isinstance(value, datetime) else str(value)
        store_id = sheet.cell(row, 13).value
        if store_id:
            key = (day, str(store_id))
            if key in existing:
                raise ValueError(f'台账已有重复 日期+店铺ID: {key}')
            existing[key] = row
    recorded_at = recorded_at or datetime.now().astimezone().isoformat(timespec='seconds')
    for store in report['stores']:
        status = store['status']
        values = [report['dateIso'], store['name'], store['platformName'],
                  float(store['sales']) if store['sales'] is not None else None, store['orders'],
                  float(store['adCost']) if store['adCost'] is not None else None,
                  float(store['roi']) if store['roi'] is not None else None,
                  status, '；'.join(filter(None, [store.get('missingReason'), store.get('adMissingReason')])), recorded_at,
                  links['html'], links.get('png'),
                  store['storeId'], store['currency'], store['orderMetric']]
        row = existing.get((report['dateIso'], store['storeId']), sheet.max_row + 1)
        for column, value in enumerate(values, 1):
            # Assign .value so a failed rerun clears old numeric cells instead of retaining them.
            sheet.cell(row, column).value = value
        for column in (4, 6, 7):
            sheet.cell(row, column).number_format = '0.00'
        sheet.cell(row, 5).number_format = '0'
    sheet.auto_filter.ref = f'A1:O{sheet.max_row}'
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    workbook.close()
    from store_ledger import update_ledgers
    update_ledgers(report)
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('report_json')
    parser.add_argument('workbook', nargs='?')
    parser.add_argument('output_dir', nargs='?')
    args = parser.parse_args()
    report = update_ledger_and_enrich(read_json(args.report_json), args.workbook, args.output_dir)
    write_json(args.report_json, report)
    print(json.dumps({'ok': True, 'reportJson': str(Path(args.report_json)),
                      'workbook': str(args.workbook or paths(report['configSnapshot'], report['dateIso'])['workbook']),
                      'rowsWritten': len(report['stores'])}, ensure_ascii=False))


if __name__ == '__main__':
    main()
